"""Strateji notunun uygulama dosyasını (XLSX) üretir.

Sayfalar: Is_Panosu, BF_Fiyat_Takvimi, Google_Kurulum, KPI_Donem
Kullanım: python3 uygulama_excel.py
Çıktı:    ../2026-10-24-uygulama.xlsx
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

BURASI = Path(__file__).resolve().parent
CIKTI = BURASI.parent / "2026-10-24-uygulama.xlsx"
BASLIK_DOLGU = PatternFill("solid", fgColor="0F1B2D")
BASLIK_FONT = Font(bold=True, color="FFFFFF")


def baslik_satiri(ws, basliklar: list[str], genislikler: list[int]) -> None:
    ws.append(basliklar)
    for i, h in enumerate(ws[1], start=0):
        h.fill = BASLIK_DOLGU
        h.font = BASLIK_FONT
        h.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[h.column_letter].width = genislikler[i]
    ws.freeze_panes = "A2"


def is_panosu(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Is_Panosu"
    baslik_satiri(ws, ["Kart", "Sütun", "Son tarih", "Sorumlu", "Not"], [58, 14, 12, 12, 50])
    kartlar = [
        ("Meta ürün gezinme kararı (kapat / test)", "Yapılacak", dt.date(2026, 10, 27), "Hasan", "Toplu kapatma 28 Ekim'e kadar"),
        ("Advantage+ creative özellik denetimi", "Yapılacak", dt.date(2026, 10, 29), "Hasan", "30 Ekim koşulları"),
        ("Meta hesap düzeyi yerleşim kontrolleri + %5 kutusu", "Yapılacak", dt.date(2026, 10, 31), "Hasan", ""),
        ("Google kanal önceliği / Display tavanı", "Yapılacak", dt.date(2026, 11, 7), "Hasan", "Kademeli açılıyor, hesapta kontrol"),
        ("BF fiyat geçmişi logu (kanal bazlı)", "Yapılacak", dt.date(2026, 10, 31), "Marka", "10 gün kuralı"),
        ("Avatar fotoğrafını repoya yükle", "Yapılacak", None, "Hasan", "pazarlama/sosyal-medya/_assets/avatar.jpg"),
        ("Super Burst uygunluğu (Reach Planner / Meta temsilcisi)", "Beklemede", None, "Hasan", "TR ülke listesi açıklanmadı"),
        ("4 karusel + video + metinler", "Biten", dt.date(2026, 10, 24), "Claude", "Avatar gelince son slayt yenilenecek"),
        ("Strateji notu (DOCX) + uygulama Excel", "Biten", dt.date(2026, 10, 24), "Claude", ""),
        ("Skill + hafıza dosyaları", "Biten", dt.date(2026, 10, 24), "Claude", ".claude/skills, CLAUDE.md"),
        ("Sabah 08:03 haber özeti görevi", "Biten", dt.date(2026, 10, 24), "Claude", "Push bildirimi"),
    ]
    for k in kartlar:
        ws.append(list(k))
    dv = DataValidation(type="list", formula1='"Yapılacak,Devam Eden,Beklemede,Biten"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"B2:B{ws.max_row + 50}")
    for c in ws["C"][1:]:
        c.number_format = "DD.MM.YYYY"


def bf_takvim(wb: Workbook) -> None:
    ws = wb.create_sheet("BF_Fiyat_Takvimi")
    baslik_satiri(ws, ["Kampanya", "İndirim başlangıcı", "Pencere başı", "Pencere sonu", "Not"], [32, 16, 14, 14, 60])
    satirlar = [
        ("Kasım boyu kampanya", dt.date(2026, 11, 1)),
        ("11.11 ısınma", dt.date(2026, 11, 9)),
        ("Black Friday haftası (erken)", dt.date(2026, 11, 20)),
        ("Black Friday haftası", dt.date(2026, 11, 23)),
        ("Black Friday günü", dt.date(2026, 11, 27)),
        ("Cyber Monday", dt.date(2026, 11, 30)),
    ]
    for i, (ad, bas) in enumerate(satirlar, start=2):
        # Pencere formülle hesaplanır: başlangıç tarihini değiştirirseniz pencere güncellenir
        ws.append([ad, bas, f"=B{i}-10", f"=B{i}-1",
                   "Bu aralıkta fiyat artırmayın; referans fiyat = aralıktaki en düşük fiyat (aynı kanal)"])
    for kol in ("B", "C", "D"):
        for c in ws[kol][1:]:
            c.number_format = "DD.MM.YYYY"


def google_kurulum(wb: Workbook) -> None:
    ws = wb.create_sheet("Google_Kurulum")
    baslik_satiri(ws, ["Faz", "Adım", "Durum", "Tarih", "Not"], [10, 70, 14, 12, 40])
    adimlar = [
        ("0", "MCC + marka hesabı, yetkiler, 2FA"),
        ("0", "Çerez yönetimi + açık rıza + aydınlatma metni (KVKK)"),
        ("0", "Consent Mode v2 (4 sinyal) GTM ile"),
        ("0", "GA4 + dönüşüm aksiyonları (satın alma birincil, değerli)"),
        ("0", "Gelişmiş dönüşümler (enhanced conversions)"),
        ("0", "Lead markası: CRM çevrimdışı dönüşüm içe aktarma"),
        ("0", "Merchant Center feed: başlık, GTIN, fiyat, stok, kargo, iade"),
        ("1", "Brand Search kampanyası"),
        ("1", "Non-brand Search (tema bazlı reklam grupları)"),
        ("1", "Performance Max (feed + varlık), marka hariç tutma"),
        ("1", "Demand Gen (yeni kitle)"),
        ("1", "Kanal önceliği: Display + iş ortakları Düşük, tavan"),
        ("1", "Hesap düzeyi negatif anahtar kelime listesi"),
        ("2", "2 hafta öğrenme; bütçe değişimi ≤ %20"),
        ("2", "30+ dönüşüm/30 gün sonrası tCPA / tROAS geçişi"),
        ("2", "Haftalık: arama terimleri, kanal raporu, varlık raporu, feed"),
        ("3", "Aylık kanal önceliği gözden geçirme"),
        ("3", "Büyük dönemlerde mevsimsellik ayarı aç / kapat"),
    ]
    for faz, adim in adimlar:
        ws.append([faz, adim, "Yapılacak", None, ""])
    dv = DataValidation(type="list", formula1='"Yapılacak,Devam Eden,Beklemede,Biten"')
    ws.add_data_validation(dv)
    dv.add(f"C2:C{ws.max_row + 50}")


def kpi_donem(wb: Workbook) -> None:
    ws = wb.create_sheet("KPI_Donem")
    baslik_satiri(ws, ["Metrik", "Önceki dönem", "Bu dönem", "Değişim %"], [34, 16, 16, 14])
    ws.append(["Dönem adı", "BF 2025 haftası", "BF 2026 haftası", ""])
    girdiler = ["Harcama (TL)", "Gösterim", "Tıklama", "Dönüşüm", "Gelir (TL)", "Lead", "Nitelikli lead (CRM)"]
    for g in girdiler:
        ws.append([g, None, None, None])
    # Satır numaraları: 3 Harcama, 4 Gösterim, 5 Tıklama, 6 Dönüşüm, 7 Gelir, 8 Lead, 9 Nitelikli lead
    hesaplar = [
        ("CPA (TL)", "=IFERROR({k}3/{k}6,\"\")"),
        ("ROAS", "=IFERROR({k}7/{k}3,\"\")"),
        ("CVR", "=IFERROR({k}6/{k}5,\"\")"),
        ("CTR", "=IFERROR({k}5/{k}4,\"\")"),
        ("Lead kalitesi (nitelikli / toplam)", "=IFERROR({k}9/{k}8,\"\")"),
    ]
    for ad, f in hesaplar:
        ws.append([ad, f.format(k="B"), f.format(k="C"), None])
    for r in range(3, ws.max_row + 1):
        ws[f"D{r}"] = f'=IFERROR((C{r}-B{r})/B{r},"")'
        ws[f"D{r}"].number_format = "0.0%"
    for r in (12, 13, 14):  # CVR, CTR, lead kalitesi yüzde
        ws[f"B{r}"].number_format = ws[f"C{r}"].number_format = "0.00%"
    for r in (3, 7, 10):
        ws[f"B{r}"].number_format = ws[f"C{r}"].number_format = "#,##0.00"
    ws[f"B11"].number_format = ws[f"C11"].number_format = "0.00"


def uret() -> Path:
    wb = Workbook()
    is_panosu(wb)
    bf_takvim(wb)
    google_kurulum(wb)
    kpi_donem(wb)
    wb.save(CIKTI)
    return CIKTI


if __name__ == "__main__":
    print(uret())
