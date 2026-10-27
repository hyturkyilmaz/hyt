"""Paid media ve e-ticaret strateji notunu (DOCX) üretir.

Kullanım: python3 strateji_notu.py
Çıktı:    ../2026-10-24-strateji-notu.docx
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor

BURASI = Path(__file__).resolve().parent
CIKTI = BURASI.parent / "2026-10-24-strateji-notu.docx"
BUGUN = dt.date(2026, 10, 24)
LACIVERT = RGBColor(0x0F, 0x1B, 0x2D)


def referans_penceresi(baslangic: dt.date) -> tuple[dt.date, dt.date]:
    """İndirim başlangıcından önceki 10 günlük referans fiyat penceresi (1.8.2026 yönetmelik değişikliği)."""
    return baslangic - dt.timedelta(days=10), baslangic - dt.timedelta(days=1)


def tr_tarih(t: dt.date) -> str:
    return t.strftime("%d.%m.%Y")


def baslik(doc: Document, metin: str, seviye: int = 1) -> None:
    h = doc.add_heading(metin, level=seviye)
    for r in h.runs:
        r.font.color.rgb = LACIVERT


def paragraf(doc: Document, metin: str, kalin: bool = False) -> None:
    p = doc.add_paragraph()
    r = p.add_run(metin)
    r.bold = kalin


def maddeler(doc: Document, liste: list[str], numarali: bool = False) -> None:
    stil = "List Number" if numarali else "List Bullet"
    for m in liste:
        doc.add_paragraph(m, style=stil)


def tablo(doc: Document, basliklar: list[str], satirlar: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(basliklar))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, b in enumerate(basliklar):
        hucre = t.rows[0].cells[i]
        hucre.text = b
        for r in hucre.paragraphs[0].runs:
            r.bold = True
    for satir in satirlar:
        hucreler = t.add_row().cells
        for i, deger in enumerate(satir):
            hucreler[i].text = deger
    doc.add_paragraph()


def uret() -> Path:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("Paid Media ve E-ticaret Strateji Notu", 0)
    paragraf(doc, f"Hasan Yasin Türkyılmaz · {tr_tarih(BUGUN)} · Sürüm 1.0")
    paragraf(
        doc,
        "Varsayım: Not, Türkiye'de faaliyet gösteren ve Meta + Google'da performans kampanyası yürüten bir "
        "e-ticaret veya lead odaklı marka için yazıldı. Bilgiler 24.10.2026 itibarıyla ikincil kaynaklardan "
        "(sektör yayınları, ajans blogları) derlendi. Bu ortamda platformların yardım sayfalarına doğrudan "
        "erişilemedi. Uygulamadan önce Ads Manager / Google Ads arayüzünde teyit edin.",
    )

    # ---------------------------------------------------------------- Yönetici özeti
    baslik(doc, "Yönetici Özeti")
    paragraf(doc, "Durum", kalin=True)
    maddeler(doc, [
        "28 Ekim: Meta, mevcut reklamlarda ürün gezinmeyi (product browsing) varsayılan açıyor. Tek tıkla toplu kapatma bu tarihe kadar.",
        "30 Ekim: Meta'nın yeni ticari koşulları. AI özelliklerinin ürettiği içerikten reklamveren sorumlu.",
        "Meta satış/lead/etkileşim/uygulama hedeflerinde reklam seti düzeyinde yerleşim hariç tutma kalkıyor. Kontrol hesap düzeyine taşındı.",
        "Google Ads: kanal önceliği (Yüksek/Düşük) ve bütçe dağılımı (yüzde tavan) kademeli açılıyor. Kasım öncesi kurulmalı.",
        "Türkiye: 1.8.2026'dan beri indirim reklamında referans fiyat = indirimden önceki 10 günün en düşük fiyatı (aynı kanal). "
        "1 Kasım'da başlayacak bir kampanyanın penceresi 22 Ekim'de açıldı.",
    ])
    paragraf(doc, "Yaptıklarım", kalin=True)
    maddeler(doc, [
        "Altı başlıkta güncel kaynak taraması (Meta, Google, Shopify, FirstPageSage, Ticaret Bakanlığı mevzuatı).",
        "Bu not, uygulama Excel'i, 4 karusel (20 slayt), 1 dikey video ve paylaşım metinleri.",
        "Skill ve hafıza dosyaları (repo: .claude/skills, CLAUDE.md). Her sabah 08:03'te çalışan haber özeti görevi.",
    ])
    paragraf(doc, "Yapacaklarım", kalin=True)
    maddeler(doc, [
        "Avatar fotoğrafın gelince 4 postun son slaytını yeniden üretmek.",
        "Platform yardım sayfalarına erişim olduğunda bu nottaki 'teyit edilmedi' maddelerini güncellemek.",
    ])
    paragraf(doc, "Senin yapacakların", kalin=True)
    maddeler(doc, [
        "Bugün-27 Ekim: Meta ürün gezinme kararını ver (Bölüm 1). Kapatacaksan 28 Ekim'den önce.",
        "30 Ekim'den önce: Advantage+ creative özelliklerini tek tek gözden geçir (Bölüm 2).",
        "Fiyat takvimini kilitle: kampanya başlangıcından önceki 10 günde fiyat artırma (Bölüm 3).",
        "Meta hesap düzeyinde yerleşim kontrollerini ve %5 kutusunu ayarla (Bölüm 5).",
        "Google'da Display ve iş ortaklarını Düşük önceliğe al, tavan koy (Bölüm 6).",
    ], numarali=True)
    paragraf(doc, "Öneri", kalin=True)
    paragraf(
        doc,
        "Kasım'a girerken otomasyonu kapatmak yerine sınırlarını çizin. Meta'da kontrol hesap ayarlarına, Google'da "
        "kanal önceliğine geçti. Bu iki ayarı doğru yapan marka, aynı bütçeyle daha az kalitesiz trafik satın alır. "
        "İndirim reklamlarında ise mevzuat riski medya riskinden büyük. Fiyat geçmişini belgelemeden kampanya açmayın.",
    )

    # ---------------------------------------------------------------- 1. Meta ürün gezinme
    baslik(doc, "1. Meta: Ürün gezinme (product browsing) 28 Ekim'de varsayılan açılıyor")
    paragraf(doc, "Ne oluyor?", kalin=True)
    paragraf(
        doc,
        "Bir kullanıcı Facebook veya Instagram'daki reklamınıza dokunduğunda, sitenize gitmeden önce Meta'nın kendi "
        "içinde barındırdığı bir ürün sayfası açılıyor. Kullanıcı ürünleri orada geziyor, isterse oradan sitenize geçiyor. "
        "Meta bu ayarı 28 Ekim 2026'da mevcut reklamlarda varsayılan olarak açıyor. Bildirimdeki toplu kapatma butonu "
        "28 Ekim'e kadar çalışıyor. Sonrasında reklam bazında Advantage+ creative / hedef (destination) ayarlarından "
        "tek tek kapatılabiliyor.",
    )
    paragraf(doc, "Meta'nın gerekçesi", kalin=True)
    paragraf(doc, "Yaklaşık 117.700 reklamverenle yapılan testte ortalama %3 dönüşüm artışı (Meta beyanı, bağımsız teyit yok).")
    paragraf(doc, "Belirsizlikler (24.10.2026 itibarıyla)", kalin=True)
    maddeler(doc, [
        "28 Ekim'den sonra açılan yeni reklamların varsayılan olarak açık başlayıp başlamayacağı yazmıyor.",
        "Hangi format, yerleşim ve hedeflerin kapsandığı belirtilmemiş.",
        "Kapatmanın tek hesaba mı tüm hesaplara mı uygulandığı belli değil.",
        "15 Ekim itibarıyla Meta yardım merkezinde resmi bir makale yoktu.",
    ])
    paragraf(doc, "Artıları ve eksileri", kalin=True)
    tablo(doc, ["Artı", "Eksi"], [
        ["Daha hızlı açılan sayfa, mobilde daha az kopuş", "Siteye gelen trafik ve remarketing kitlesi küçülebilir"],
        ["Meta'nın beyan ettiği ~%3 dönüşüm artışı", "GA4 / site analitiğinde atıf zayıflar, Meta raporuna bağımlılık artar"],
        ["Kataloğu güçlü markalarda keşif artabilir", "Siteye özel unsurlar (kampanya bandı, paket, kupon) görünmeyebilir"],
        ["", "Birinci taraf veri (e-posta, üyelik) toplama fırsatı azalır"],
    ])
    paragraf(doc, "Karar kuralı", kalin=True)
    maddeler(doc, [
        "Kapatın: Birinci taraf veri toplamak, remarketing havuzunu büyütmek veya sitede özel teklif göstermek önceliğinizse.",
        "Kapatın: Lead kampanyalarında form kalitesini CRM ile ölçüyorsanız (Meta içi akış veri kalitesini değiştirebilir).",
        "Test edin: Katalog satışı yapan, sitesi mobilde yavaş olan markalarda kampanyaların yarısında açık, yarısında kapalı 14 gün.",
        "Test metriği: Site + Meta toplam satın alma, blended ROAS, yeni müşteri oranı. Sadece Meta'nın kendi raporuna bakmayın.",
    ])
    paragraf(
        doc,
        "Not: Kararsızsanız 28 Ekim'den önce toplu kapatıp sonra kontrollü test açmak, açık bırakıp sonra reklam reklam "
        "kapatmaktan daha az iş çıkarır. Bu yüzden önerim önce kapatmak, sonra test etmek.",
    )

    # ---------------------------------------------------------------- 2. Meta 30 Ekim
    baslik(doc, "2. Meta: 30 Ekim ticari koşulları ve AI sorumluluğu")
    maddeler(doc, [
        "Koşullara 'AI Özellikleri' bölümü ekleniyor. Meta'nın AI özelliklerinin ürettiği veya değiştirdiği içerikten reklamveren sorumlu.",
        "AI'ın hata yapabileceği açıkça yazıyor; risk reklamverende.",
        "AI tarafından eklenen filigran veya içerik doğrulama metaverisi değiştirilemez ve silinemez.",
        "Tetikleyici olay: Bazı markaların reklamları, hesap güncellemesiyle açılan özellikler yüzünden habersizce 'güzelleştirildi' ve yayından çekildi.",
        "Meta'ya göre bu bir 'düzenleme', reklamveren Advantage+ özelliklerini tek tek kapatabilir.",
    ])
    paragraf(doc, "Yapılacaklar", kalin=True)
    maddeler(doc, [
        "Tüm aktif reklamlarda Advantage+ creative alt özelliklerini listeleyin (görsel iyileştirme, metin varyasyonu, arka plan üretimi, müzik vb.).",
        "Regüle alanlarda (sağlık, gıda takviyesi, finans, kozmetik iddiaları) metin ve görsel üreten özellikleri kapatın.",
        "Açık bıraktıklarınız için haftalık ekran görüntüsü arşivi tutun. Reklam Kurulu incelemesinde kanıt olur.",
        "Türkiye: 1.8.2026 değişikliğiyle AI ile üretilen reklamlar için açıklama yükümlülüğü gündemde. Resmi metni hukuk ekibiyle kontrol edin.",
    ])

    # ---------------------------------------------------------------- 3. Black Friday
    baslik(doc, "3. Black Friday / Efsane Cuma strateji notu (Türkiye mevzuatıyla)")
    paragraf(doc, "3.1 Kural özeti (Ticari Reklam ve Haksız Ticari Uygulamalar Yönetmeliği, değişiklik RG 33297, yürürlük 1.8.2026)", kalin=True)
    maddeler(doc, [
        "İndirim öncesi fiyat: indirimin başladığı tarihten önceki 10 gün içinde uygulanan en düşük fiyat (önceden 30 gündü).",
        "Sadece indirimin yapıldığı satış kanalındaki fiyat dikkate alınır. Site ile pazaryeri ayrı değerlendirilir.",
        "Çabuk bozulan ürünler ve hizmetlerde eski kural sürüyor: indirimden hemen önceki fiyat.",
        "Kolay erişilebilen sadakat programı indirimleri de indirim kurallarına tabi.",
        "Hedefli reklam kavramı ilk kez yönetmelikte tanımlandı. AI ile üretilen reklamlar ve influencer tanıtımları kapsamda.",
        "Denetim: Reklam Kurulu 8 Ekim toplantısında 161 dosyada 136 durdurma, toplam 17.176.744 TL ceza. 2026'nın ilk 9 ayında 249 milyon TL'yi aşan ceza. "
        "Ekim kararlarında 10 gün kuralına uymama ve sadakat kartı kampanyaları öne çıktı.",
    ])
    paragraf(doc, "3.2 Fiyat takvimi (kodla hesaplandı)", kalin=True)
    satirlar = []
    for bas, ad in [
        (dt.date(2026, 11, 1), "Kasım boyu kampanya"),
        (dt.date(2026, 11, 9), "11.11 ısınma"),
        (dt.date(2026, 11, 20), "Black Friday haftası (erken)"),
        (dt.date(2026, 11, 23), "Black Friday haftası"),
        (dt.date(2026, 11, 27), "Black Friday günü"),
    ]:
        p1, p2 = referans_penceresi(bas)
        satirlar.append([ad, tr_tarih(bas), f"{tr_tarih(p1)} – {tr_tarih(p2)}"])
    tablo(doc, ["Kampanya", "İndirim başlangıcı", "Referans fiyat penceresi (fiyat artırma!)"], satirlar)
    paragraf(doc, "Takvim: Black Friday 27.11.2026 (Cuma), Cyber Monday 30.11.2026 (Pazartesi). Bazı kaynaklarda geçen 1 Aralık hatalı.")
    paragraf(doc, "3.3 Yapın / yapmayın", kalin=True)
    tablo(doc, ["Yapın", "Yapmayın"], [
        ["Referans penceresi boyunca her ürünün günlük fiyatını kanal bazında kaydedin (ekran görüntüsü + veritabanı logu)",
         "Pencere içinde fiyat artırıp sonra 'indirim' göstermek"],
        ["'…'e varan' ifadesinde en yüksek oranın gerçekten uygulandığı ürünleri belgeleyin ve kampanya süresince stokta tutun",
         "Tek ürüne uygulanan oranı tüm kategoriye genellemek"],
        ["Kart, üyelik, sepette indirim gibi koşulları reklamın kendisinde açık yazın", "Koşulu sadece dipnotta veya ödeme adımında göstermek"],
        ["Sınırlı stok veya süre varsa gerçek rakamı yazın", "Sahte geri sayım sayacı, sürekli yenilenen 'son saatler' ifadesi"],
        ["Influencer işbirliklerinde reklam etiketini (#reklam / işbirliği) baştan koyun", "Etiketi yorumlara veya hashtag yığınının sonuna saklamak"],
        ["AI ile üretilen görselleri iç onaydan geçirin, gerekiyorsa belirtin", "Ürünü olduğundan farklı gösteren AI düzenlemesi"],
    ])
    paragraf(doc, "3.4 Medya planı (başlangıç hipotezi, hesabınızın verisiyle ayarlanmalı)", kalin=True)
    tablo(doc, ["Faz", "Tarih", "Meta", "Google"], [
        ["Isınma / veri", "1–15 Kasım", "Video ve katalog ile kitle büyütme, e-posta/üyelik toplama", "Brand + non-brand Search, Demand Gen ile ilgi toplama"],
        ["Hazırlık", "16–22 Kasım", "Remarketing kitleleri, kreatif testleri bitirilir", "PMax varlıkları, feed promosyonları, kanal önceliği ayarı"],
        ["Zirve", "23–30 Kasım", "Gün bazlı bütçe artışı; uygunsa 27 Kasım için Super Burst", "Mevsimsellik ayarı (seasonality adjustment), tROAS'ı geçici gevşetme"],
        ["Uzatma", "1–7 Aralık", "Sepet terk ve yeni müşteri dönüşü", "Bütçeyi kademeli normale çekme, mevsimsellik ayarını kapatma"],
    ])
    paragraf(doc, "3.5 KPI seti ve dönem karşılaştırması", kalin=True)
    paragraf(
        doc,
        "Harcama, CPA, ROAS, CVR, lead kalitesi (CRM'de satışa dönen lead oranı). Karşılaştırma: BF 2025 haftası vs BF 2026 haftası "
        "ve Kasım 2026 vs Ekim 2026. Excel dosyasındaki 'KPI_Donem' sayfası formüllü şablon olarak hazır.",
    )

    # ---------------------------------------------------------------- 4. Super Burst / First Impression
    baslik(doc, "4. Meta Ads Super Burst ve First Impression Ads")
    paragraf(doc, "4.1 Ne bunlar?", kalin=True)
    tablo(doc, ["", "Ads Super Burst", "First Impression Ads"], [
        ["Duyuru", "15 Ekim 2026, Brand Building Summit", "15 Ekim 2026, aynı etkinlik"],
        ["Ne", "3–24 saatlik rezerve kampanya penceresi, Facebook + Instagram", "Facebook akışını açan kişinin gördüğü ilk reklam yeri"],
        ["Satın alma", "Reach Planner'da rezervasyon; erişim ve fiyat tahmini, tahmine bağlı garantili teslim", "Açıklanmadı (sabit CPM rezervasyon iddiası teyitsiz)"],
        ["Rekabet", "Ek ücretle kategori münhasırlığı: aynı pencerede rakip kategori reklam alamaz", "Günde tek marka"],
        ["Meta'nın test verisi (ABD)", "Medyan 34 milyon kişi/uçuş (ABD online yetişkinlerin %14'ü), reklam hatırlamada +2,9 puan, niyette +0,6 puan", "Yayınlanmadı"],
        ["Erişilebilirlik", "Ekim'de kademeli genel kullanıma geçiş; ülke listesi açıklanmadı", "Sadece ABD'de erken test"],
    ])
    paragraf(doc, "4.2 Avantajımız ne olur?", kalin=True)
    maddeler(doc, [
        "Zaman sıkıştırma: Normalde günlere yayılan erişimi birkaç saate toplar. 'Herkes aynı anda duydu' etkisi yaratır. Lansman, sınırlı stok, canlı yayın, maç anı gibi anlarda değerli.",
        "Öngörülebilirlik: Açık artırma yerine rezervasyon. Fiyat ve erişim önceden belli. Black Friday gibi CPM'in fırladığı günlerde bütçe planı netleşir.",
        "Rakip blokajı: Kategori münhasırlığıyla o saatlerde rakibin reklamı aynı kitleye çıkmaz. Pazaryeri ve perakendede ciddi fark yaratabilir.",
        "Performansı besler: Kısa sürede yaratılan bilinirlik markalı arama, doğrudan trafik ve remarketing havuzunu büyütür. Sonraki günlerin performans kampanyaları daha ucuza çalışır (hipotez, ölçülmeli).",
        "First Impression Ads: Günün ilk reklam yeri. Dikkat en yüksekken tek marka. Büyük lansman veya kriz/duyuru iletişimi için uygun.",
    ])
    paragraf(doc, "4.3 Ne amaçla kullanılır, ne zaman kullanılmaz?", kalin=True)
    tablo(doc, ["Kullanın", "Kullanmayın"], [
        ["Ürün lansmanı, koleksiyon açılışı", "Günlük satış hedefi olan sürekli performans kampanyası"],
        ["Black Friday / 11.11 açılış saatleri", "Kitle 200.000 kişinin altındaysa (rezervasyon modelinin bilinen alt sınırı)"],
        ["Maç, final, ödül töreni gibi kültürel anlar", "Kreatif tek ve test edilmemişse (birkaç saatte hata telafi edilemez)"],
        ["Rakibin büyük kampanya gününe karşı savunma", "Ölçüm altyapısı (brand lift, markalı arama takibi) yoksa"],
    ])
    paragraf(doc, "4.4 Stratejiye yerleştirme", kalin=True)
    maddeler(doc, [
        "Bütçe: Toplam medya bütçesinin %5–10'unu yılda 3–4 'an'a ayırın. Geri kalan performans tarafında kalır (öneri, hesap verisiyle ayarlanmalı).",
        "Kurgu: Super Burst günü + sonraki 3–7 gün güçlendirilmiş retargeting ve markalı Search bütçesi.",
        "Ölçüm: Meta Brand Lift (uygunsa), Google'da markalı arama hacmi, doğrudan trafik, yeni müşteri oranı. Önceki 4 haftanın aynı günleriyle kıyas.",
        "Türkiye: Ülke listesi açıklanmadı. Meta temsilcinizden ya da Reach Planner'da hesabınızın uygunluğunu kontrol edin. Yoksa mevcut Reservation (eski Reach & Frequency) ile benzer kurgu kurulabilir.",
        "Açık sorular: Fiyat seviyesi, minimum harcama, münhasırlık ücreti. Rakam verilmeden bütçe taahhüdü yapmayın.",
    ])

    # ---------------------------------------------------------------- 5. Yerleşim hariç tutma
    baslik(doc, "5. Meta yerleşim hariç tutma kalktı: Para boşa gitmesin")
    paragraf(
        doc,
        "Satış, lead, etkileşim ve uygulama hedeflerinde reklam seti düzeyinde manuel yerleşim kalkıyor (Bilinirlik ve Trafik koruyor). "
        "Kaynaklar arasında Etkileşim'in kapsamı konusunda çelişki var; hesabınızda kontrol edin. Değişiklik Ağustos sonundan beri kademeli.",
    )
    maddeler(doc, [
        "Hesap düzeyi engel: Reklam Ayarları > Hesap Kontrolleri > Yerleşim kontrolleri. Hedef bazında yerleşim kapatılabiliyor. Bu dört hedefte tek gerçek engel bu.",
        "'%5'e kadar harca' kutusu: 'Performansı artıracaksa her hariç yerleşime bütçenin %5'ine kadar harca' varsayılan açık. Tam engel için kapatın.",
        "Değer kuralları: Zayıf yerleşimde teklifi %90'a kadar düşürür. Engellemez. En fazla 10 kural, kural başına 2 kriter. Meta, toplam sonuç başı maliyetin artabileceğini uyarıyor.",
        "Yerleşime özel kreatif: Reels'e 9:16 video, akışa 4:5, Stories'e güvenli alanlı tasarım. Kötü uyan kreatif, kötü yerleşimden daha pahalıya patlar.",
        "Marka güvenliği: Engelleme listeleri ve envanter filtresi aynı ayar alanında.",
        "Hesap mimarisi: Hesap düzeyi ayar tüm kampanyalara işlediği için farklı marka güvenliği ihtiyacı olan işleri ayrı reklam hesabına ayırın.",
        "Ölçüm: Haftalık yerleşim kırılımı. Harcama, CPA, ROAS, CVR ve CRM'den lead kalitesi yan yana. Lead formlarında Audience Network ve Messenger gelen kutusu kaynaklı lead'lerin satışa dönüşünü ayrıca izleyin.",
        "Eşik kuralı (öneri): Bir yerleşim 2 hafta üst üste hesap ortalamasının 1,5 katı CPA veya yarısı kadar lead kalitesi verirse hesap düzeyinde kapatın.",
    ], numarali=True)

    # ---------------------------------------------------------------- 6. Google
    baslik(doc, "6. Google Ads: Kanal önceliği, bütçe dağılımı, API v26, Universal Cart")
    paragraf(doc, "6.1 Kanal önceliği ve bütçe dağılımı (duyuru 14.10.2026, kademeli)", kalin=True)
    maddeler(doc, [
        "Öncelik: Her kanal Yüksek veya Düşük; varsayılan Yüksek. Google, hedef CPA/ROAS içinde önce Yüksek kanalları fonlar, kalanı Düşük kanallara verir.",
        "Bütçe dağılımı: Düşük kanallara günlük ortalama bütçenin yüzdesi kadar tavan. %0 kanalı kapatır. Hesap, belirtilen bütçe üzerinden yapılır, gerçekleşen harcama üzerinden değil. Kesin limit değil, bazı günler aşabilir.",
        "Kapsam: Search, Shopping, Performance Max, Demand Gen. Demand Gen'de öncelik var, bütçe dağılımı yok.",
        "PMax / Demand Gen sınırı: Search, YouTube, Maps, Gmail, Discover 'Google kanalları' olarak tek blok. Arama ağı iş ortakları ve Display ayrı kademelendirilebilir.",
        "Seviyeler: Hesap varsayılanı, kampanya ve paylaşılan bütçe düzeyinde ayarlanabilir. Kampanya ayarı hesap varsayılanını ezer.",
        "Kampanya toplam bütçesi: Bir kaynağa göre Search, Shopping ve PMax'te kampanya toplam bütçesi kullanılıyorsa ayar uygulanmıyor (teyit edilmedi).",
        "Son tarih: Display Expansion kullanan Search kampanyaları Şubat 2027'de geçiyor. Mevcut davranış için Google'ın önerisi Display Düşük + %50.",
    ])
    paragraf(doc, "6.2 Nasıl avantaja çeviririz?", kalin=True)
    tablo(doc, ["Senaryo", "Ayar", "Neden"], [
        ["Küçük bütçeli marka", "Display %0, iş ortakları Düşük %5", "Sınırlı bütçe kaliteli aramaya gitsin"],
        ["Black Friday", "Google kanalları Yüksek; Display + iş ortakları Düşük %10–15", "Zirvede ucuz envantere kaçışı sınırla"],
        ["Yeni kanal testi", "Test kanalı Düşük %5–10, 2–4 hafta", "Riski sınırlı, öğrenme verisi alınır"],
        ["Paylaşılan bütçeli portföy", "Paylaşılan bütçe düzeyinde tek kural", "Kampanyalar arası tutarlılık"],
        ["Search + Display Expansion", "Şubat 2027 öncesi Display Düşük + %50 ile başla, sonra veriye göre düşür", "Geçişte performans kırılmasın"],
    ])
    paragraf(doc, "6.3 Google Ads API v26: ayda iki sürüm", kalin=True)
    maddeler(doc, [
        "v26'dan itibaren ayda iki sürüm. Yılda yine dört büyük sürüm. Ara sürümler sadece geriye uyumlu yenilik ekler.",
        "Sunset tarihleri değişmedi; daha sık geçiş gerekmiyor. Blog ve canlı yayın duyuruları üç ayda ikiye iniyor.",
        "2027 taslak takvim: v27 Ocak, v28 Nisan, v29 Temmuz, v30 Ekim (tahmini).",
        "Bizim için anlamı: Script ve entegrasyonlarda yeni alanlar (ör. kanal önceliği) daha hızlı gelebilir. Kendi araçlarınızda (ör. WordPress içerik otomasyonu, raporlama) sürüm izlemeyi aylık değil iki haftalık yapın.",
    ])
    paragraf(doc, "6.4 Universal Cart (ABD, 15.10.2026)", kalin=True)
    maddeler(doc, [
        "Search (AI Mode dahil) ve Gemini'de, farklı mağazalardan ürünleri tek sepette tutan alışveriş özelliği. Google Pay ile sepette ödeme veya ürünleri mağazanın sepetine aktarma.",
        "Satıcı kayıtlı satıcı olarak kalır; sipariş onayı ve destek mağazada.",
        "İlk ortaklar: Nike, Sephora, Target, Ulta, Walmart, Wayfair ve çok sayıda Shopify mağazası. YouTube ve Gmail sonra; Avustralya ve Kanada gelecek yıl (Merchant Center UCP entegrasyonu).",
        "Türkiye için: Şimdilik doğrudan etkisi yok. ABD'ye ihracat yapan markalar Merchant Center feed kalitesini ve UCP hazırlığını bugünden ele almalı.",
    ])
    paragraf(doc, "6.5 Bir markada sıfırdan Google Ads kurulumu", kalin=True)
    paragraf(doc, "Faz 0 — Altyapı (1. hafta)", kalin=True)
    maddeler(doc, [
        "Hesap yapısı: MCC altında marka hesabı, faturalandırma, kullanıcı yetkileri (en az yetki ilkesi, 2FA zorunlu).",
        "KVKK: Çerez yönetim platformu, açık rıza metni, aydınlatma metni. Consent Mode v2 (ad_storage, ad_user_data, ad_personalization, analytics_storage). Yurt dışına veri aktarımı için hukuki dayanak.",
        "Ölçüm: GA4 + Google Tag Manager. Dönüşüm aksiyonları: satın alma (birincil, değerli), sepete ekleme ve ödeme başlatma (ikincil). Gelişmiş dönüşümler (enhanced conversions) aydınlatma metniyle.",
        "Lead markaları: CRM'den çevrimdışı dönüşüm içe aktarma (nitelikli lead, satış). Lead kalitesi ölçülmeden tCPA kurmayın.",
        "Merchant Center: Feed (başlık, açıklama, GTIN, marka, fiyat, stok, kargo, iade politikası), promosyonlar, ürün yorumları.",
    ], numarali=True)
    paragraf(doc, "Faz 1 — Kampanya mimarisi (2. hafta)", kalin=True)
    tablo(doc, ["Kampanya", "Amaç", "Teklif (başlangıç)", "Kanal önceliği", "Bütçe payı (hipotez)"], [
        ["Brand Search", "Markayı koru, ucuz dönüşüm", "Maks. dönüşüm veya hedef gösterim payı", "Search Yüksek, iş ortakları Düşük %0–5", "%10–15"],
        ["Non-brand Search (tema bazlı)", "Talep yakalama", "Maks. dönüşüm → yeterli veriyle tCPA/tROAS", "Display %0, iş ortakları Düşük %5", "%30–40"],
        ["Performance Max (feed + varlık)", "Ölçekli satış", "Maks. dönüşüm değeri → tROAS", "Google kanalları Yüksek; Display + iş ortakları Düşük %10", "%35–45"],
        ["Demand Gen", "Talep yaratma, yeni kitle", "Maks. dönüşüm / tıklama", "Öncelik ayarlanır (bütçe dağılımı yok)", "%10–15"],
    ])
    maddeler(doc, [
        "PMax'te marka hariç tutma (brand exclusions) açın; brand Search ile çakışmasın.",
        "Kitle sinyalleri: müşteri listesi (KVKK rızalı), site ziyaretçisi, özel segment (rakip ve kategori aramaları).",
        "Negatif anahtar kelime listeleri: iş, ücretsiz, ikinci el, nasıl yapılır vb. Hesap düzeyi liste.",
    ])
    paragraf(doc, "Faz 2 — Öğrenme ve teklif geçişi (3–8. hafta)", kalin=True)
    maddeler(doc, [
        "İlk 2 hafta öğrenme dönemine dokunmayın; günlük bütçeyi %20'den fazla değiştirmeyin (yaygın uygulama).",
        "Kampanya başına son 30 günde 30+ dönüşüm oluşunca hedef CPA/ROAS'a geçin. Hedefi gerçekleşenin %10–20 gevşeğinden başlatın.",
        "Haftalık: arama terimleri, kanal performansı raporu, varlık raporu, feed hataları.",
    ])
    paragraf(doc, "Faz 3 — Ölçek ve kontrol (2. aydan itibaren)", kalin=True)
    maddeler(doc, [
        "Kanal önceliği kararlarını kanal performansı raporuna göre güncelleyin (aylık).",
        "Büyük dönemlerde mevsimsellik ayarı (seasonality adjustment) kullanın, dönem bitince kapatın.",
        "Rapor: Harcama, CPA, ROAS, CVR, lead kalitesi; haftalık ve aylık dönem karşılaştırması.",
    ])

    # ---------------------------------------------------------------- 7. Shopify
    baslik(doc, "7. Shopify ve agentic commerce")
    maddeler(doc, [
        "Agentic Storefronts (Mart 2026 duyurusu): Shopify Catalog üzerinden ChatGPT, Copilot, Google AI Mode ve Gemini'ye ürün akışı. Meta tarafı kademeli. Uygun mağazalarda varsayılan açık.",
        "28 Eylül 2026 (TechCrunch): Tarayıcı tabanlı AI ajanları Shopify ödeme sayfasını tamamlayabiliyor.",
        "'Shop MCP' (13 Ekim) adı tek kaynakta geçiyor. Shopify'ın resmi tarafında Storefront MCP ve AI Toolkit var. Ad teyit edilmedi.",
        "Protokol: Universal Commerce Protocol (Google + Shopify, Ocak 2026). REST, MCP, AP2 ve A2A üzerinden çalışabiliyor.",
        "Amazon: Dış ajanları engelliyor (Meta Muse vakası). Kendi ajanı içeride.",
        "Ülke durumu: AI üzerinden ödeme büyük ölçüde ABD mağazalarında; uluslararası yayılım 2026 2. yarısı hedefli. Katalog görünürlüğü ülke sınırı içermiyor.",
    ])
    paragraf(doc, "Türkiye'deki markalar için aksiyon", kalin=True)
    maddeler(doc, [
        "Shopify Admin > Satış kanalları: AI kanallarının açık olup olmadığını kontrol edin, bilinçli karar verin.",
        "Ürün verisi: başlık (marka + model + temel özellik), yapılandırılmış özellikler, GTIN, stok, kargo süresi, iade koşulu.",
        "Yapılandırılmış veri (schema.org Product, Offer, Review) sitenizde eksiksiz olsun. Shopify dışı altyapılarda (WooCommerce, ikas, Ticimax vb.) aynı ilke geçerli.",
        "Analitik: chatgpt.com, gemini.google.com, copilot.microsoft.com yönlendirmelerini GA4'te ayrı kanal grubu yapın.",
        "Hukuk: Ajan üzerinden satışta mesafeli satış ve ön bilgilendirme yükümlülüklerinin nasıl karşılandığını kontrol edin.",
    ])

    # ---------------------------------------------------------------- 8. FirstPageSage
    baslik(doc, "8. FirstPageSage verisi: Araştırma ile satın alma arasındaki fark")
    maddeler(doc, [
        "Kapsam: 13 Temmuz – 16 Ekim 2026, 4.936 kişilik ABD online alışverişçi paneli + 43 markanın sipariş/yönlendirme verisi.",
        "%61,4 satın alma araştırmasında AI kullanıyor; %6,8 ödemeyi ajana tamamlatmış.",
        "Ajanın tamamladığı sipariş payı %0,27 → %0,68 (2026 2. → 3. çeyrek).",
        "AI destekli oturumların %41,3'ü ChatGPT'de; ajanla tamamlanan satın almaların %9,4'ü ChatGPT'de; Alexa for Shopping + Meta Muse %61,5.",
        "Ajan ödeme ortaklığı olan mağazalarda tamamlanma 1,4–3,8 kat.",
        "Sınırlılık: Tek kaynak, ticari bir SEO ajansının paneli, ABD verisi. Bağımsız teyit yok. Farklı çalışmalar farklı güven oranları veriyor (ör. YouGov %14 güven).",
        "Yorum: Bugün bütçe ve emek, AI araştırma aşamasında görünür olmaya (ürün verisi, yorum, karşılaştırma içeriği) gitmeli. Ajan ödemesi izlenecek trend.",
    ])

    # ---------------------------------------------------------------- 9. Güven tablosu
    baslik(doc, "9. Güven düzeyleri")
    tablo(doc, ["Konu", "Güven", "Neden"], [
        ["Meta ürün gezinme 28 Ekim", "Orta", "Çok sayıda ikincil kaynak, resmi yardım sayfası yok"],
        ["Meta 30 Ekim koşulları", "Yüksek", "Koşul sayfası bannerı ve Adweek dahil çok kaynak"],
        ["Yerleşim hariç tutma kaldırılması", "Yüksek (kapsam orta)", "Yardım sayfası alıntılanıyor; Etkileşim kapsamında çelişki"],
        ["Super Burst / First Impression", "Yüksek (detay düşük)", "Duyuru net; fiyat, ülke, satın alma yolu açıklanmadı"],
        ["Google kanal önceliği", "Yüksek", "Birden çok PPC yayını, aynı tarih ve detay"],
        ["API v26 iki sürüm", "Yüksek", "Birden çok kaynak tutarlı"],
        ["Universal Cart", "Yüksek", "Digital Commerce 360, TechTimes ve diğerleri"],
        ["Shopify agentic", "Orta", "Ülke ve kanal durumu kaynaklar arasında farklı"],
        ["FirstPageSage rakamları", "Orta-düşük", "Tek kaynak, bağımsız teyit yok"],
        ["TR yönetmelik 10 gün kuralı", "Yüksek", "Bakanlık duyurusu + çok sayıda hukuk bülteni; resmi metin RG 33297"],
    ])

    # ---------------------------------------------------------------- Kaynaklar
    baslik(doc, "Kaynaklar")
    paragraf(doc, "Akademik kaynak bulunmamaktadır; konular çok yeni olduğu için tüm kaynaklar sektör yayını, ajans blogu ve resmi duyurulardır.")
    for k in KAYNAKLAR:
        doc.add_paragraph(k, style="List Bullet")

    doc.save(CIKTI)
    return CIKTI


KAYNAKLAR = [
    "Adweek. (2026). Advertisers pull Meta ads that had unwanted AI glow-ups. https://www.adweek.com/media/advertisers-pull-meta-ads-that-had-unwanted-ai-glow-ups/",
    "Brain & Code. (2026). Shopify lets AI agents check out. Amazon locks them out. https://brainandcode.com/shopify-agent-checkout-vs-amazon/",
    "CNN Türk. (2026, Ekim). Reklam Kurulu'ndan 249 milyon lirayı aşan ceza: İndirim reklamlarına sıkı denetim. https://www.cnnturk.com/ekonomi/reklam-kurulundan-249-milyon-lirayi-asan-ceza-indirim-reklamlarina-siki-denetim-3480168",
    "Common Thread Co. (2026). Google Ads now lets you choose which channels get priority budget. https://commonthreadco.com/blogs/coachs-corner/google-ads-now-lets-you-choose-which-channels-get-priority-budget-what-every-ecommerce-brand-must-do-before-bfcm",
    "Common Thread Co. (2026). Meta launches Ads Super Burst and First Impression Ads. https://commonthreadco.com/blogs/coachs-corner/meta-ads-super-burst-first-impression-ads-bfcm-2026",
    "Craftshift. (2026). Which Shopify AI channels work outside the US? https://craftshift.com/shopify-ai-channels-outside-us/",
    "Digital Commerce 360. (2026, 16 Ekim). Google launches Universal Cart for purchases across retailers. https://www.digitalcommerce360.com/2026/10/16/google-universal-cart-launch-features/",
    "FirstPageSage. (2026). Agentic commerce statistics: October 2026. https://firstpagesage.com/seo-blog/agentic-commerce-statistics",
    "Gromerce. (2026). Meta removed placement exclusions from your ad sets. https://gromerce.com/en/blog/meta-placement-exclusions-removed-value-rules-august-2026",
    "Kurşun & Partners. (2026). İndirimli satış reklamlarında son 10 gün en düşük fiyat kuralı. https://kursunpartners.com/tr/hukuk-bulteni/gelisme/indirimli-satis-reklamlarinda-son-10-gun-en-dusuk-fiyat-kurali/",
    "PPC Land. (2026). Meta makes advertisers solely responsible for its AI ad edits from Oct 30. https://ppc.land/meta-makes-advertisers-solely-responsible-for-its-ai-ad-edits-from-oct-30/",
    "PPC Land. (2026). Meta sells brands 3 to 24 hour ad windows that shut out category rivals. https://ppc.land/meta-sells-brands-3-to-24-hour-ad-windows-that-shut-out-category-rivals/",
    "PPC Land. (2026). Google Ads API moves to two releases a month starting with v26. https://ppc.land/google-ads-api-moves-to-two-releases-a-month-starting-with-v26/",
    "Relevant Audience. (2026). Google Ads channel prioritization and budget caps. https://www.relevantaudience.com/google-ads-en/google-ads-channel-prioritization-budget-allocation/",
    "Relevant Audience. (2026). Meta's replacements for ad set placement controls. https://www.relevantaudience.com/meta/meta-placement-controls-alternatives-documented/",
    "Reklam Özdenetim Kurulu. (2026). Ticari Reklam ve Haksız Ticari Uygulamalar Yönetmeliği değişikliği. https://www.rok.org.tr/haberler/ticari-reklam-ve-haksiz-ticari-uygulamalar-yonetmeligi-degisikligi/",
    "Search Engine Land. (2026). Meta introduces new brand advertising formats across Facebook and Instagram. https://searchengineland.com/meta-introduces-new-brand-advertising-formats-across-facebook-and-instagram-496608",
    "Search Engine Roundtable. (2026). New Google Ads controls: Channel prioritization & budget allocation. https://www.seroundtable.com/google-ads-channel-prioritization-budget-allocation-42297.html",
    "Social Media Today. (2026). Meta offers updated ad placement control guidance. https://www.socialmediatoday.com/news/meta-offers-updated-ad-placement-control-guidance/832791/",
    "T.C. Ticaret Bakanlığı. (2026). Ticari reklam ve haksız ticari uygulamalar yönetmeliğinde yapılan değişikliklerle tüketicilerin korunması güçlendiriliyor. https://ticaret.gov.tr/haberler/ticaret-bakanligi-tarafindan-ticari-reklam-ve-haksiz-ticari-uygulamalar-yonetmeliginde-yapilan-degisikliklerle-tuketicilerin-aldatici-reklam-ve-ticari-uygulamalara-karsi-korunmasi-guclendiriliyor",
    "TechCrunch. (2026, 28 Eylül). Shopify opens checkout to browser-based AI agents. https://techcrunch.com/2026/09/28/shopify-opens-checkout-to-browser-based-ai-agents/",
    "The Keyword. (2026). Meta is making product browsing the default for existing ads. https://www.thekeyword.co/news/meta-product-browsing-default",
    "TheOptimizer. (2026). Meta ads placement control in 2026: How to actually block placements. https://theoptimizer.io/blog/meta-ads-placement-control-in-2026-how-to-actually-block-placements-its-not-as-simple-anymore",
    "Ticaret Gazetesi. (2026, 16 Ekim). Reklam Kurulu, aldatıcı reklamlara karşı 17,2 milyon TL idari para cezası uyguladı. https://ticaretgazetesi.com.tr/2026/10/16/reklam-kurulu-aldatici-reklamlara-karsi-172-milyon-tl-idari-para-cezasi-uyguladi/",
]


if __name__ == "__main__":
    print(uret())
