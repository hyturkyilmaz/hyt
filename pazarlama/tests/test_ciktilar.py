"""Pazarlama çıktılarının doğrulama testleri.

Çalıştırma: python3 -m unittest discover -s pazarlama/tests -v   (repo kökünden)
"""
from __future__ import annotations

import datetime as dt
import json
import subprocess
import sys
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from PIL import Image

KOK = Path(__file__).resolve().parents[1]
SOSYAL = KOK / "sosyal-medya"
sys.path.insert(0, str(KOK / "_kaynak"))

import strateji_notu  # noqa: E402

ZORUNLU_CTA = "danışmanlığı için iletişime geçin"
IMZA_ISIM = "Hasan Yasin Türkyılmaz"


class PostKurallari(unittest.TestCase):
    def setUp(self):
        self.veri = json.loads((SOSYAL / "_kaynak" / "posts.json").read_text(encoding="utf-8"))

    def test_imza_ve_cta_metni(self):
        imza = self.veri["imza"]
        self.assertEqual(imza["isim"], IMZA_ISIM)
        self.assertIn(ZORUNLU_CTA, imza["cta"])
        for alan in ("AI", "pazarlama", "reklam", "e-ticaret"):
            self.assertIn(alan, imza["cta"])

    def test_her_post_cta_ile_biter_ve_3_6_slayt(self):
        for post in self.veri["postlar"]:
            with self.subTest(post=post["id"]):
                self.assertEqual(post["slaytlar"][-1]["tip"], "cta")
                self.assertTrue(3 <= len(post["slaytlar"]) <= 6)

    def test_klasor_adlari_ascii(self):
        for post in self.veri["postlar"]:
            self.assertTrue(post["id"].isascii(), post["id"])

    def test_icerik_slaytlarinda_kaynak_var(self):
        for post in self.veri["postlar"]:
            for s in post["slaytlar"]:
                if s["tip"] == "icerik":
                    self.assertTrue(s.get("kaynak"), f"{post['id']}: kaynaksız içerik slaytı")


class GorselCiktilar(unittest.TestCase):
    def test_png_sayisi_ve_boyut(self):
        veri = json.loads((SOSYAL / "_kaynak" / "posts.json").read_text(encoding="utf-8"))
        for post in veri["postlar"]:
            for i in range(1, len(post["slaytlar"]) + 1):
                yol = SOSYAL / post["id"] / f"slayt-{i}.png"
                with self.subTest(yol=str(yol)):
                    self.assertTrue(yol.exists())
                    with Image.open(yol) as img:
                        self.assertEqual(img.size, (1080, 1350))

    def test_video_dikey(self):
        video = SOSYAL / "02-google-kanal-onceligi" / "google-kanal-onceligi-9x16.mp4"
        self.assertTrue(video.exists())
        cikti = subprocess.run(
            ["ffprobe", "-v", "error", "-select_streams", "v:0", "-show_entries", "stream=width,height",
             "-of", "csv=p=0", str(video)], capture_output=True, text=True, check=True).stdout.strip()
        self.assertEqual(cikti, "1080,1920")


class RenderGuvenlik(unittest.TestCase):
    def test_html_kacis(self):
        # render.mjs içindeki kaçış fonksiyonu HTML enjeksiyonunu engellemeli
        kod = (SOSYAL / "_kaynak" / "render.mjs").read_text(encoding="utf-8")
        self.assertIn("&lt;", kod)
        self.assertIn("&quot;", kod)
        js = ("const kac=" + kod.split("const kac =", 1)[1].split(";\n", 1)[0] +
              ";console.log(kac('<script>alert(\"x\")</script>'))")
        sonuc = subprocess.run(["node", "-e", js], capture_output=True, text=True, check=True).stdout.strip()
        self.assertNotIn("<script>", sonuc)
        self.assertIn("&lt;script&gt;", sonuc)

    def test_render_dis_ag_kullanmaz(self):
        kod = (SOSYAL / "_kaynak" / "render.mjs").read_text(encoding="utf-8")
        self.assertNotIn("http://", kod)
        self.assertNotIn("https://", kod)


class StratejiNotu(unittest.TestCase):
    def test_referans_penceresi(self):
        self.assertEqual(strateji_notu.referans_penceresi(dt.date(2026, 11, 27)),
                         (dt.date(2026, 11, 17), dt.date(2026, 11, 26)))
        self.assertEqual(strateji_notu.referans_penceresi(dt.date(2026, 11, 1)),
                         (dt.date(2026, 10, 22), dt.date(2026, 10, 31)))

    def test_bf_ve_cyber_monday_gunleri(self):
        self.assertEqual(dt.date(2026, 11, 27).weekday(), 4)  # Cuma
        self.assertEqual(dt.date(2026, 11, 30).weekday(), 0)  # Pazartesi

    def test_docx_bolumleri(self):
        doc = Document(KOK / "2026-10-24-strateji-notu.docx")
        basliklar = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        metin = "\n".join(p.text for p in doc.paragraphs)
        for beklenen in ("Yönetici Özeti", "Kaynaklar", "Super Burst", "sıfırdan Google Ads"):
            self.assertTrue(any(beklenen in b for b in basliklar) or beklenen in metin, beklenen)
        for etiket in ("Durum", "Yaptıklarım", "Yapacaklarım", "Senin yapacakların", "Öneri"):
            self.assertIn(etiket, metin)


class UygulamaExcel(unittest.TestCase):
    def setUp(self):
        self.wb = load_workbook(KOK / "2026-10-24-uygulama.xlsx")

    def test_sayfalar(self):
        self.assertEqual(self.wb.sheetnames, ["Is_Panosu", "BF_Fiyat_Takvimi", "Google_Kurulum", "KPI_Donem"])

    def test_formuller(self):
        ws = self.wb["BF_Fiyat_Takvimi"]
        self.assertEqual(ws["C2"].value, "=B2-10")
        self.assertEqual(ws["D2"].value, "=B2-1")
        kpi = self.wb["KPI_Donem"]
        self.assertEqual(kpi["A10"].value, "CPA (TL)")
        self.assertEqual(kpi["B10"].value, '=IFERROR(B3/B6,"")')
        self.assertEqual(kpi["A11"].value, "ROAS")
        self.assertEqual(kpi["B11"].value, '=IFERROR(B7/B3,"")')

    def test_pano_sutunlari(self):
        ws = self.wb["Is_Panosu"]
        izinli = {"Yapılacak", "Devam Eden", "Beklemede", "Biten"}
        for r in ws.iter_rows(min_row=2, values_only=True):
            self.assertIn(r[1], izinli)


if __name__ == "__main__":
    unittest.main()
